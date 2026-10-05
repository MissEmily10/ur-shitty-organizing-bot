"""Текст из документов: PDF, Word, Excel, TXT/MD/CSV."""

import csv
import io
import re
from dataclasses import dataclass, field

# Сколько страниц скана отдаём модели как картинки: каждая страница тратит кредиты
SCAN_PAGES = 6
# Сколько строк таблицы Excel берём с каждого листа
SHEET_ROWS = 300


class Unsupported(Exception):
    pass


@dataclass
class Extracted:
    text: str
    kind: str  # «PDF», «Word», «Таблица», «Текст»
    scans: list[bytes] = field(default_factory=list)  # страницы скана, если в PDF нет текстового слоя


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1251"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _reflow(text: str) -> str:
    """PDF рвёт строки по ширине страницы: склеиваем строки внутри абзаца, абзацы разделяем пустой строкой."""
    text = re.sub(r"-\n(?=\w)", "", text)  # перенос слова по слогам
    paragraphs = re.split(r"\n\s*\n", text)
    return "\n\n".join(re.sub(r"\s*\n\s*", " ", p).strip() for p in paragraphs if p.strip())


def _md_table(rows: list[list]) -> str:
    rows = [["" if v is None else str(v).replace("|", "/").replace("\n", " ") for v in row] for row in rows]
    rows = [r for r in rows if any(cell.strip() for cell in r)]
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    lines = ["| " + " | ".join(rows[0]) + " |", "|" + "---|" * width]
    lines += ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return "\n".join(lines)


def _jpeg(data: bytes) -> bytes:
    """Страницу скана — в JPEG: легче для модели и для загрузки."""
    from PIL import Image

    image = Image.open(io.BytesIO(data))
    image.thumbnail((2000, 2000))
    out = io.BytesIO()
    image.convert("RGB").save(out, "JPEG", quality=85)
    return out.getvalue()


def _pdf(data: bytes) -> Extracted:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = [page.extract_text() or "" for page in reader.pages]
    text = "\n\n".join(_reflow(p) for p in pages if p.strip())
    if len(text) >= 40 * max(1, min(len(pages), 3)):
        return Extracted(text, "PDF")
    # Почти нет текста — это скан: берём картинки страниц и отдаём их модели, как фото
    scans = []
    for page in reader.pages[:SCAN_PAGES]:
        images = list(page.images)
        if images:
            scans.append(_jpeg(max(images, key=lambda im: len(im.data)).data))
    if not scans:
        raise Unsupported("В PDF нет ни текста, ни картинок страниц.")
    return Extracted(text, "PDF", scans)


def _docx(data: bytes) -> Extracted:
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts = []
    for p in doc.paragraphs:
        if not p.text.strip():
            continue
        style = (p.style.name or "").lower() if p.style is not None else ""
        level = re.search(r"heading (\d)", style)
        parts.append(("#" * min(int(level.group(1)), 3) + " " if level else "") + p.text.strip())
    for table in doc.tables:
        parts.append(_md_table([[cell.text for cell in row.cells] for row in table.rows]))
    return Extracted("\n\n".join(p for p in parts if p), "Word")


def _xlsx(data: bytes) -> Extracted:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        rows = [list(row) for _, row in zip(range(SHEET_ROWS), ws.iter_rows(values_only=True))]
        if table := _md_table(rows):
            parts.append(f"## {ws.title}\n{table}")
    return Extracted("\n\n".join(parts), "Таблица")


def extract(data: bytes, filename: str, mime: str | None) -> Extracted:
    name = filename.lower()
    if name.endswith(".pdf") or mime == "application/pdf":
        return _pdf(data)
    if name.endswith(".docx"):
        return _docx(data)
    if name.endswith((".xlsx", ".xlsm")):
        return _xlsx(data)
    if name.endswith(".csv"):
        rows = list(csv.reader(io.StringIO(_decode(data)), delimiter=";" if b";" in data[:2000] else ","))
        return Extracted(_md_table(rows[:SHEET_ROWS]), "Таблица")
    if name.endswith((".txt", ".md")) or (mime or "").startswith("text/"):
        return Extracted(_decode(data).strip(), "Текст")
    if name.endswith((".doc", ".xls")):
        raise Unsupported("Старые форматы .doc и .xls не поддерживаются: пересохраните файл как .docx или .xlsx.")
    raise Unsupported("Этот формат пока не поддерживается. Подходят PDF, Word (.docx), Excel (.xlsx), CSV, TXT, MD.")
